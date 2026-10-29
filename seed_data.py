import os
import shutil
import openpyxl
from database import init_db, add_expense, get_connection

APP_DIR = os.path.dirname(os.path.abspath(__file__))
SOURCE_DIR = '/Users/code_red/.gemini/antigravity/scratch/bill_merger'

PAYMENTS_UPLOAD_DIR = os.path.join(APP_DIR, 'data', 'uploads', 'payments')
BILLS_UPLOAD_DIR = os.path.join(APP_DIR, 'data', 'uploads', 'bills')

os.makedirs(PAYMENTS_UPLOAD_DIR, exist_ok=True)
os.makedirs(BILLS_UPLOAD_DIR, exist_ok=True)

def seed():
    init_db()
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM expenses")
    count = c.fetchone()[0]
    conn.close()

    if count > 0:
        print(f"Database already contains {count} records. Skipping seed.")
        return

    excel_path = os.path.join(SOURCE_DIR, 'Bills excel (1) (1).xlsx')
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active

    pp_dir = os.path.join(SOURCE_DIR, 'Phone pe screenshot')
    bill_dir = os.path.join(SOURCE_DIR, 'softcopy_pages')

    pp_files = sorted([f for f in os.listdir(pp_dir) if f.lower().endswith('.jpg')])
    bill_files = sorted([f for f in os.listdir(bill_dir) if f.lower().endswith('.jpg') and not f.startswith('opt_')])

    # Verified mapping
    MAPPING = {
        1: {'pp': [1, 2], 'bill': [32]},
        2: {'pp': [3], 'bill': [38]},
        3: {'pp': [4, 5], 'bill': []},
        4: {'pp': [6], 'bill': [1]},
        5: {'pp': [7], 'bill': [40]},
        6: {'pp': [8], 'bill': [33]},
        7: {'pp': [9], 'bill': [2]},
        8: {'pp': [10], 'bill': [3]},
        9: {'pp': [11], 'bill': [8]},
        10: {'pp': [12], 'bill': []},
        11: {'pp': [13], 'bill': [4]},
        12: {'pp': [14], 'bill': [5]},
        13: {'pp': [15], 'bill': []},
        14: {'pp': [16], 'bill': [9]},
        15: {'pp': [17], 'bill': [6]},
        16: {'pp': [18], 'bill': []},
        17: {'pp': [19], 'bill': [7]},
        18: {'pp': [20], 'bill': [10]},
        19: {'pp': [21], 'bill': [11]},
        20: {'pp': [22], 'bill': []},
        21: {'pp': [23], 'bill': []},
        22: {'pp': [24], 'bill': [12]},
        23: {'pp': [25], 'bill': [13]},
        24: {'pp': [26], 'bill': [14]},
        25: {'pp': [27], 'bill': [19]},
        26: {'pp': [28], 'bill': [15]},
        27: {'pp': [29, 30], 'bill': [16]},
        28: {'pp': [31], 'bill': [17]},
        29: {'pp': [32, 33], 'bill': [20]},
        30: {'pp': [34], 'bill': []},
        31: {'pp': [35], 'bill': [18]},
        32: {'pp': [36], 'bill': [21]},
        33: {'pp': [37], 'bill': []},
        34: {'pp': [38, 39], 'bill': [22]},
        35: {'pp': [40], 'bill': []},
        36: {'pp': [41], 'bill': [39]},
        37: {'pp': [42, 43], 'bill': [23]},
        38: {'pp': [44], 'bill': [24]},
        39: {'pp': [45], 'bill': [25]},
        40: {'pp': [46], 'bill': [26]},
        41: {'pp': [47], 'bill': []},
        42: {'pp': [48], 'bill': [34]},
        43: {'pp': [49], 'bill': [27]},
        44: {'pp': [50], 'bill_custom': os.path.join(bill_dir, 'redbus_2666_ticket.jpg')},
        45: {'pp': [51], 'bill': [28]},
        46: {'pp': [52], 'bill': []},
        47: {'pp': [53, 54, 55], 'bill': [30]},
        48: {'pp': [56], 'bill': []},
        49: {'pp': [57], 'bill': [35]},
        50: {'pp': [58], 'bill': [36]},
        51: {'pp': [59], 'bill': [29]},
        52: {'pp': [60], 'bill': [37]},
        53: {'pp': [61], 'bill': []},
        54: {'pp': [62], 'bill': [31]}, # KOT slip
        55: {'pp': [63, 64], 'bill': []},
        56: {'pp': [65], 'bill_custom': os.path.join(bill_dir, 'uber_149_oct18.jpg')}
    }

    # Month conversion for standardized ISO date sorting
    # e.g., '6 -Oct' -> '2026-10-06'
    def to_iso_date(raw_date):
        d_str = str(raw_date).replace('\xa0', ' ').strip()
        parts = d_str.split('-')
        if len(parts) == 2:
            day = parts[0].strip().zfill(2)
            mon = parts[1].strip()
            if 'oct' in mon.lower():
                return f"2026-10-{day}"
            elif 'nov' in mon.lower():
                return f"2026-11-{day}"
        return d_str

    print("Seeding 56 historical records from October tour...")
    for r in range(2, sheet.max_row+1):
        s_no = sheet.cell(r, 1).value
        if s_no is None:
            continue
        s_no = int(s_no)
        raw_date = str(sheet.cell(r, 2).value or '').strip()
        iso_date = to_iso_date(raw_date)
        desc = str(sheet.cell(r, 3).value or '').strip().replace('\xa0', ' ')
        cat = str(sheet.cell(r, 4).value or '').strip().replace('\xa0', ' ')
        amt = str(sheet.cell(r, 5).value or '').strip().replace('\xa0', ' ')
        paid_by = str(sheet.cell(r, 6).value or 'Shiraz').strip()
        remark = str(sheet.cell(r, 8).value or '').strip().replace('\xa0', ' ').replace('₹', 'Rs. ')

        if s_no == 44:
            remark = "Return travel ticket to hyd, RedBus ticket copy attached (PNR #1821860)."
        elif s_no == 56:
            remark = "Uber ride receipt attached (Rs. 148.48)."

        info = MAPPING[s_no]

        # Copy payment screenshots to uploads folder
        saved_payments = []
        for p_idx in info['pp']:
            src = os.path.join(pp_dir, pp_files[p_idx - 1])
            dst_name = f"payment_{s_no}_{os.path.basename(src)}"
            dst = os.path.join(PAYMENTS_UPLOAD_DIR, dst_name)
            shutil.copy2(src, dst)
            saved_payments.append(dst)

        # Copy bill screenshots to uploads folder
        saved_bills = []
        if 'bill_custom' in info and os.path.exists(info['bill_custom']):
            src = info['bill_custom']
            dst_name = f"bill_{s_no}_{os.path.basename(src)}"
            dst = os.path.join(BILLS_UPLOAD_DIR, dst_name)
            shutil.copy2(src, dst)
            saved_bills.append(dst)
        else:
            for b_idx in info.get('bill', []):
                src = os.path.join(bill_dir, bill_files[b_idx - 1])
                dst_name = f"bill_{s_no}_{os.path.basename(src)}"
                dst = os.path.join(BILLS_UPLOAD_DIR, dst_name)
                shutil.copy2(src, dst)
                saved_bills.append(dst)

        add_expense(
            date=iso_date,
            description=desc,
            category=cat,
            amount_expression=amt,
            paid_by=paid_by,
            remark=remark,
            payments=saved_payments,
            bills=saved_bills
        )

    print("Successfully seeded all 56 records!")

if __name__ == '__main__':
    seed()
