import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel_report(expenses, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expense Report"

    # Header styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    double_bottom_border = Border(
        top=Side(style='thin', color='1E3A8A'),
        bottom=Side(style='double', color='1E3A8A')
    )

    headers = [
        "S. No", "Date", "Description", "Category", 
        "Amount(INR)", "Paid By", "Sheet Page number", "Remark"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 7] else "left", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    # Insert Data rows
    for idx, e in enumerate(expenses):
        row_num = idx + 2
        s_no = idx + 1
        page_num = idx + 1 # 1-to-1 match with PDF voucher page
        amt_val = e.get('amount_expression', e.get('amount_total', 0))

        row_data = [
            s_no,
            str(e.get('date', '')),
            str(e.get('description', '')),
            str(e.get('category', '')),
            amt_val,
            str(e.get('paid_by', 'Shiraz') or 'Shiraz'),
            page_num,
            str(e.get('remark', '') or '')
        ]
        ws.append(row_data)

        # Style data row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [1, 2, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row_num].height = 20

    # Total row
    last_row = len(expenses) + 2
    total_claim = sum(float(e.get('amount_total', 0) or 0) for e in expenses)

    ws.cell(row=last_row, column=3, value="Total Paid:").font = bold_font
    ws.cell(row=last_row, column=3).alignment = Alignment(horizontal="right", vertical="center")

    total_cell = ws.cell(row=last_row, column=5, value=round(total_claim, 2))
    total_cell.font = Font(name="Calibri", size=11, bold=True, color="047857")
    total_cell.number_format = "#,##0.00"
    total_cell.alignment = Alignment(horizontal="right", vertical="center")

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=last_row, column=col_idx)
        cell.fill = total_fill
        cell.border = double_bottom_border

    ws.row_dimensions[last_row].height = 24

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    ws.column_dimensions['C'].width = 32 # Description
    ws.column_dimensions['H'].width = 45 # Remark

    wb.save(output_path)
    return output_path
